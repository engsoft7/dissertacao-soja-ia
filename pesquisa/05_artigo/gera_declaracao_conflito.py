# -*- coding: utf-8 -*-
"""
Declaração de conflito de interesse em arquivo separado, como o formulário pede.

O manuscrito já traz a declaração na seção própria, mas o Editorial Manager
exige um .docx à parte — e exige que seja Word, não PDF nem texto colado. São
dois lugares com a mesma frase, e é por isso que este gerador existe: a frase
mora aqui e em gera_manuscrito_cea.py, e se divergirem o autor assina duas
declarações diferentes no mesmo processo.

O texto é o mesmo, palavra por palavra, do que o manuscrito imprime.
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.shared import Cm, Pt

RAIZ = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(RAIZ, 'saida')
os.makedirs(SAIDA, exist_ok=True)

F = 'Times New Roman'
TITULO = 'Declaration of competing interest'
DECLARACAO = (
    'The author declares no known competing financial interests or personal '
    'relationships that could have appeared to influence the work reported in '
    'this paper.'
)
ARTIGO = ('Value repetition in official crop statistics: a diagnostic of '
          'target-variable quality and its effect on the evaluation of machine '
          'learning yield models')

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(2.5)
s.left_margin = s.right_margin = Cm(2.5)
d.styles['Normal'].font.name = F
d.styles['Normal'].font.size = Pt(11)


def p(t='', neg=False, align=AL.JUSTIFY, dep=10, tam=11):
    par = d.add_paragraph()
    par.alignment = align
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.space_after = Pt(dep)
    if t:
        r = par.add_run(t)
        r.font.name, r.font.size, r.font.bold = F, Pt(tam), neg
    return par


p(TITULO, neg=True, align=AL.LEFT, dep=14, tam=12)
p('Manuscript: ' + ARTIGO, align=AL.LEFT, dep=6)
p('Author: Maycon Lima dos Santos (ORCID 0009-0005-9424-6048)',
  align=AL.LEFT, dep=16)
p(DECLARACAO, dep=16)
p('Maycon Lima dos Santos', align=AL.LEFT, dep=0)
p('Graduate Program in Applied Computing, Universidade Federal do Pará',
  align=AL.LEFT, dep=0)
p('14 September 2026', align=AL.LEFT, dep=0)

d.save(os.path.join(SAIDA, 'Declaration_of_competing_interests.docx'))
print('gerada')
